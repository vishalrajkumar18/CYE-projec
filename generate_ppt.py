from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()

    # --- Utility functions for styling ---
    def set_title_style(title_shape):
        title_shape.text_frame.paragraphs[0].font.name = 'Arial'
        title_shape.text_frame.paragraphs[0].font.size = Pt(40)
        title_shape.text_frame.paragraphs[0].font.bold = True
        title_shape.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)

    def add_bullets(body_shape, points):
        tf = body_shape.text_frame
        tf.clear()
        for point in points:
            p = tf.add_paragraph()
            p.text = point
            p.font.name = 'Arial'
            p.font.size = Pt(24)
            p.space_after = Pt(14)

    # Slide 1: Title
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    title = slide.shapes.title
    subtitle = slide.placeholders[1]
    title.text = "Usage-Based Preventive Maintenance Planner"
    subtitle.text = "Project Analysis & Overview\nFor High-Value Diagnostic Equipment"
    
    title.text_frame.paragraphs[0].font.color.rgb = RGBColor(0, 51, 102)
    title.text_frame.paragraphs[0].font.bold = True
    subtitle.text_frame.paragraphs[0].font.color.rgb = RGBColor(102, 102, 102)

    # Slide 2: Problem Statement
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "The Problem: Calendar-Based Maintenance"
    set_title_style(title)
    points = [
        "Traditional scheduling relies on fixed calendar intervals.",
        "Ignores real-world factors: equipment utilization, age, and fault history.",
        "Results in unplanned equipment downtime.",
        "Maintenance is often performed too early on low-risk machines.",
        "High-risk machines are serviced too late, jeopardizing continuity."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 3: Proposed Solution
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Proposed Solution"
    set_title_style(title)
    points = [
        "A data-driven, usage-based preventive maintenance planner.",
        "Focuses on shared high-value medical equipment (MRI, CT scanners).",
        "Dynamically calculates risk based on multifaceted metrics.",
        "Schedules maintenance to minimize disruption to patient bookings.",
        "Ensures operational continuity and reduces unexpected downtime."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 4: Key Features
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Key Features & Capabilities"
    set_title_style(title)
    points = [
        "Interpretable Risk-Scoring Engine: Configurable by administrators.",
        "Feasible Maintenance Planner: Balances hard and soft constraints.",
        "Synthetic Data Generation: Simulates real-world equipment operations.",
        "Role-Based Access Control (RBAC): Admin, Manager, Technician, Ops.",
        "Override Workflow & Audit Logging: Complete traceability."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 5: System Architecture
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "System Architecture & Stack"
    set_title_style(title)
    points = [
        "Frontend: Streamlit (Python) for rapid, interactive UI development.",
        "Backend: Python ecosystem (Pandas, Numpy, scikit-learn).",
        "Database: SQLite for lightweight, reliable data persistence.",
        "Visualization: Plotly for interactive data and risk dashboards.",
        "Quality Assurance: Pytest for unit and integration testing."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 6: Risk Scoring Engine
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Risk Scoring Engine"
    set_title_style(title)
    points = [
        "Calculates an overall risk score using a weighted sum:",
        "• 30% Utilisation: Wear and tear from heavy usage.",
        "• 20% Age: Older machines pose higher failure risks.",
        "• 20% Fault History: Frequency of recent operational faults.",
        "• 15% Service History: Time elapsed since last maintenance.",
        "• 15% Criticality: Operational importance of the specific device."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 7: Constraint-Based Planning
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Constraint-Based Planning"
    set_title_style(title)
    points = [
        "Hard Constraints (Non-negotiable):",
        "• Technician availability must be confirmed.",
        "• Must not conflict with existing patient bookings.",
        "• Must fit within the operational window of the center.",
        "Soft Constraints (Optimization):",
        "• Prefers periods of historically low utilization.",
        "• Aims to minimize overall operational disruption."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 8: Role-Based Workflows
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Role-Based Workflows"
    set_title_style(title)
    points = [
        "Admin: Manages users, configures risk weights, system settings.",
        "Maintenance Manager: Reviews recommendations, approves schedules.",
        "• Has explicit override capabilities documented via Audit Logs.",
        "Technician: Views assigned maintenance tasks, updates statuses.",
        "Operations Staff: Views overall schedule to plan patient bookings."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 9: Evaluation & Impact
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Evaluation & Business Impact"
    set_title_style(title)
    points = [
        "Target Metric: >20% reduction in unplanned downtime.",
        "Methodology: Comparing the usage-based model against a baseline calendar-only schedule.",
        "Impact: Improved equipment availability translates to higher revenue and better patient care.",
        "Dashboards: Provides real-time metrics comparing downtime strategies."
    ]
    add_bullets(slide.placeholders[1], points)

    # Slide 10: Future Enhancements
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    title = slide.shapes.title
    title.text = "Future Enhancements"
    set_title_style(title)
    points = [
        "IoT Telemetry Integration: Real-time sensor data for fault detection.",
        "Time-Series Forecasting: AI models to predict precise failure windows.",
        "Multi-Center Deployment: Scaling the platform across hospital networks.",
        "Regulatory Compliance: Integrating manufacturer maintenance requirements.",
        "Advanced Analytics: Deep learning for spare parts inventory optimization."
    ]
    add_bullets(slide.placeholders[1], points)

    # Save the presentation
    prs.save('Maintenance_Planner_Analysis.pptx')
    print("Presentation generated successfully at Maintenance_Planner_Analysis.pptx")

if __name__ == "__main__":
    create_presentation()
